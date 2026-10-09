from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path(r'C:\Users\augus\.workspace\act\git-ide\outputs\posicionamento-okestre')
d=Document(); s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(2); s.bottom_margin=Cm(2); s.left_margin=Cm(2.1); s.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.space_after=Pt(7); d.styles['Normal'].paragraph_format.line_spacing=1.1
d.styles['Title'].font.size=Pt(27); d.styles['Title'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',17),('Heading 2',12)]:
 d.styles[name].font.size=Pt(size); d.styles[name].paragraph_format.space_before=Pt(12); d.styles[name].paragraph_format.space_after=Pt(6)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(head,rows,widths):
 t=d.add_table(rows=1, cols=len(head)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,head):
  c.text=txt; sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8F1EE'); c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs:r.bold=True
 for row in rows:
  cells=t.add_row().cells
  for c,txt in zip(cells,row):c.text=txt
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c,w in zip(row.cells,widths):
   c.width=Cm(w)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(6);pp.paragraph_format.space_before=Pt(6)
    for r in pp.runs:r.font.size=Pt(10)
 return t

d.add_heading('Posicionamento da Okestre IA',0)
p('Direção de marca e proposta de valor • Outubro de 2026')
p('A Okestre IA deve se posicionar como uma plataforma de execução supervisionada de trabalho de negócio, documentação e tecnologia. Sua proposta é transformar objetivos e demandas em planos, documentos, processos e implementações verificáveis, com agentes de IA executando etapas e pessoas definindo objetivos, aprovando ações relevantes e validando resultados.')
p('Este posicionamento orienta produto, comunicação e oferta comercial. A direção central é aproximar o valor entregue do trabalho concluído, com escopo e critérios de aceite claros.')
h('Tese de posicionamento')
p('A unidade de valor da Okestre IA é a entrega aceita. A interface organiza o trabalho; a IA executa as etapas autorizadas; o responsável acompanha evidências e decide sobre a conclusão. O cliente deve conseguir entender o que foi solicitado, o que foi feito, o que foi validado e o que ainda depende de intervenção.')
h('Categoria e definição')
p('Categoria proposta: plataforma de execução de negócio, documentação e tecnologia com IA supervisionada. O conceito Service as Software expressa o modelo de entrega: atividades de um serviço são executadas por software, sob responsabilidade e supervisão humanas.')
p('Declaração de posicionamento: para equipes responsáveis por planejar, estruturar e executar iniciativas de negócio e tecnologia, a Okestre IA propõe um ambiente que conecta contexto do projeto, execução por agentes e validação humana, conduzindo demandas de negócio, documentação e implementação até entregas com evidências e critérios de aceite.')
h('Da ferramenta à entrega')
table(['Dimensão','Software as a Service','Service as Software'],[
 ('Papel do software','Oferecer recursos para operar','Executar etapas do trabalho'),
 ('Papel humano','Realizar o trabalho na ferramenta','Definir, supervisionar e validar'),
 ('Valor percebido','Acesso e produtividade de uso','Trabalho concluído e aceito'),
 ('Evidência de valor','Uso dos recursos disponíveis','Entrega e validação do resultado')],[3.1,6.8,6.9])
p('A comparação define a direção da oferta. A forma de cobrança precisa considerar também escopo, consumo, exceções e responsabilidade pela entrega.')

page('Público e proposta de valor')
h('Público prioritário proposto')
p('O foco são empresas de serviços, consultorias, software houses e equipes internas que precisam conectar negócio, documentação e execução. Os usuários incluem gestores, analistas de negócio, responsáveis por produto e projetos e profissionais técnicos. O comprador é quem responde pela capacidade de entrega e pela qualidade dos resultados dessas iniciativas.')
p('O melhor ponto de entrada é um fluxo recorrente, com acesso controlado aos recursos do projeto, critérios de conclusão observáveis e uma pessoa responsável pela revisão. Esse recorte permite demonstrar valor antes de ampliar a autonomia.')
h('Problema que a marca deve assumir')
p('Entre uma demanda e uma entrega há diagnóstico, decisões de negócio, definição de escopo, documentação, execução e acompanhamento. Quando essas etapas ficam dispersas, a pessoa precisa reconstruir o contexto, transportar informações e conferir manualmente se o pedido foi atendido. A Okestre IA deve assumir a organização e a execução desse percurso dentro de um escopo definido.')
h('Proposta de valor')
p('Delegar trabalho de negócio, documentação e tecnologia com contexto e acompanhar uma entrega verificável. O benefício esperado é reduzir o esforço de coordenação e execução repetitiva, mantendo visibilidade sobre as mudanças e responsabilidade humana pelas decisões. Ganhos de prazo e custo devem ser demonstrados em uso, sem percentuais antecipados.')
h('Diferenciação a construir')
table(['Alternativa','Ênfase da Okestre IA'],[
 ('Assistente de conversa','Executar ações no projeto e apresentar evidências do resultado, além de responder.'),
 ('IDE e copiloto','Conectar definição de negócio, documentação e implementação em um fluxo de entrega.'),
 ('Automação de tarefas','Interpretar a demanda dentro de limites definidos e lidar com exceções por supervisão.'),
 ('Serviço profissional tradicional','Tornar partes da análise, documentação e execução repetíveis por software, com aceite explícito.')],[4.5,12.3])
h('Princípios da experiência')
p('O usuário começa pelo resultado desejado. A plataforma informa o escopo, executa as ações permitidas e apresenta o que mudou. Dúvidas, bloqueios e falhas aparecem no fluxo. Uma tarefa só é concluída quando satisfaz os critérios combinados; encerrar a conversa não equivale a entregar o trabalho.')

page('Negócio e documentação no centro da entrega')
h('Estruturação de negócio')
p('A Okestre IA deve apoiar a transformação de uma necessidade em uma iniciativa executável: organizar o problema, os objetivos, o público envolvido, a proposta de valor, o escopo e as prioridades. Essas entregas orientam o trabalho seguinte e podem gerar valor por si mesmas, mesmo quando não há desenvolvimento de software.')
h('Documentação como entrega e contexto')
p('A documentação é uma frente de trabalho própria e também a base que conecta as demais etapas. Briefings, análises, propostas, requisitos, planos, procedimentos e relatórios registram decisões e orientam a execução. Sua qualidade depende de clareza, consistência, rastreabilidade e revisão pelo responsável.')
p('O posicionamento deve contemplar criação, revisão, organização e atualização de documentos. Uma mudança de objetivo ou escopo precisa ser refletida nos documentos e nas tarefas relacionadas, evitando que a execução siga orientações desatualizadas.')
h('Processos e planejamento')
p('A plataforma deve conectar a definição do negócio ao planejamento operacional: mapear etapas, responsabilidades, dependências, estimativas e critérios de aceite. A IA pode preparar e executar atividades delimitadas; as pessoas validam premissas e tomam as decisões que exigem julgamento e responsabilidade.')
h('Três frentes conectadas')
table(['Frente','Entregas possíveis','Validação esperada'],[
 ('Negócio','Diagnóstico, escopo, proposta e prioridades','Coerência com o objetivo e decisão do responsável'),
 ('Documentação','Requisitos, planos, procedimentos e relatórios','Clareza, consistência e revisão do conteúdo'),
 ('Tecnologia','Implementações, correções e configuração','Comportamento, testes e resultado observado')],[3.0,7.0,6.8])
h('Exemplo de jornada integrada')
p('Uma empresa quer organizar um novo serviço. O trabalho começa pela definição do público, do problema atendido e da oferta. Em seguida, são elaborados o processo operacional, a proposta e os requisitos. O plano de execução distribui atividades e critérios de aceite. Se houver necessidade de tecnologia, a implementação utiliza esse contexto; se a entrega for apenas de negócio ou documentação, o fluxo se encerra com sua validação.')
p('O valor está na continuidade entre intenção, decisão, documentação e execução. O catálogo comercial deve tornar explícito quais dessas entregas já estão disponíveis e validadas.')

page('Oferta e modelo de entrega')
h('Unidade de trabalho')
p('Cada entrega deve ter uma demanda identificada, contexto disponível, recursos autorizados, limites de execução e critérios de aceite. Ao concluir, a Okestre IA apresenta os artefatos produzidos, as alterações realizadas, as verificações executadas e eventuais pendências.')
h('Ciclo de execução supervisionada')
p('1. Definir a demanda e o resultado esperado.\n2. Reunir contexto e confirmar o escopo operacional.\n3. Executar as etapas autorizadas com registro das ações.\n4. Validar o resultado e apresentar evidências.\n5. Obter aceite ou retornar para ajustes.')
h('Ofertas iniciais candidatas')
table(['Entrega','Critério de aceite sugerido'],[
 ('Estruturação de negócio','Objetivo, problema, escopo e prioridades documentados e aprovados.'),
 ('Documentação e processos','Documentos coerentes, etapas e responsáveis definidos e revisão concluída.'),
 ('Planejamento de projeto','Requisitos, tarefas, estimativas e critérios de aceite consistentes.'),
 ('Implementação técnica','Comportamento ou correção validado nos cenários e ambiente acordados.')],[5.3,11.5])
p('Essas ofertas devem entrar no catálogo conforme cada fluxo for validado. Uma entrega pode terminar com bloqueio documentado; nesse caso, seu estado e tratamento comercial precisam ser claros.')
h('Direção comercial')
p('A proposta comercial deve dar destaque ao escopo das entregas e à responsabilidade assumida. Um modelo possível combina capacidade mensal de execução com unidades de trabalho aceitas. A cobrança por entrega é uma hipótese a validar, especialmente quando complexidade, infraestrutura e dependências externas variam.')
p('Antes de fixar preços, definir o que constitui uma entrega, quais revisões estão incluídas, como mudanças de escopo são tratadas e quem aprova o resultado. Consumo de modelos é um custo operacional; não substitui a definição de valor para o cliente.')
h('Indicadores para comprovar valor')
p('Acompanhar tempo entre demanda e aceite, taxa de aceite na primeira revisão, retrabalho, tempo de supervisão humana e custo por entrega aceita. Medir também falhas e intervenções. Comparar tarefas de complexidade semelhante antes de afirmar ganhos de produtividade.')

page('Comunicação e evolução do produto')
h('Mensagem principal')
p('A Okestre IA conecta negócio, documentação e tecnologia para transformar demandas em entregas executadas por IA, com supervisão humana e resultados verificáveis.')
h('Apresentação institucional')
p('A Okestre IA está sendo desenvolvida para conectar objetivos de negócio à elaboração de documentos e à execução operacional e técnica por agentes de IA. A proposta reúne contexto, planejamento, processos, ferramentas e validação em um fluxo supervisionado, no qual pessoas definem prioridades, acompanham mudanças e aprovam entregas. Seu foco é o trabalho concluído com evidências.')
h('Mensagem para página comercial')
p('Título sugerido: Da demanda à entrega com IA supervisionada.\nTexto de apoio: Estruture o negócio, organize a documentação e acompanhe a execução com critérios claros de entrega.\nChamada para ação: Escolha um fluxo para validar com a Okestre IA.')
h('Como expressar a promessa')
p('Usar verbos concretos: executar, revisar, testar, registrar e validar. Explicar quais tarefas o produto atende e onde a supervisão entra. Evitar prometer autonomia total, ausência de erros, substituição integral da equipe ou redução de custo sem medição. A promessa deve acompanhar a capacidade demonstrada em cada fluxo.')
h('Base atual e direção de evolução')
p('O escopo da Okestre IA abrange negócio e documentação, além da execução técnica. A base analisada inclui documentos, requisitos, tarefas, estimativas e recursos de projeto, com ferramentas de arquivos, Git, QA e infraestrutura e um agente próprio. Sua confiabilidade ainda precisa ser validada por fluxo. Foi planejada a substituição do motor atual pelo Agents SDK; essa substituição não foi implementada nesta etapa.')
p('A evolução deve priorizar ferramentas confiáveis, sessões persistidas, estados de execução claros, aprovações e evidências de conclusão. O SDK é uma decisão técnica de suporte à proposta. O argumento comercial é a capacidade de concluir trabalho sob supervisão, comprovada no produto.')
h('Condição para sustentar o posicionamento')
p('Começar com um conjunto restrito de entregas que funcione de ponta a ponta. Ampliar o catálogo após validar execução, recuperação de falhas e aceite. A marca ganha credibilidade quando cada entrega anunciada corresponde a um fluxo demonstrável e quando limites e pendências são comunicados com precisão.')
p('Diretriz de marca: IA executa. Pessoas supervisionam. O resultado é validado.')
d.core_properties.title='Posicionamento da Okestre IA';d.core_properties.subject='Direção de marca e proposta de valor';d.core_properties.author='Okestre IA'
for root in [d._element,d.styles.element]:
 for e in root.xpath('.//w:pBdr'): e.getparent().remove(e)
d.save(out/'Posicionamento da Okestre IA.docx')
print(out/'Posicionamento da Okestre IA.docx')
